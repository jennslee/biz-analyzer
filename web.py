#!/usr/bin/env python3
"""
사업구상 분석 웹서버
uvicorn web:app --host 0.0.0.0 --port 8200 --reload
"""

import asyncio, io, json, os, re, sqlite3
from datetime import datetime

# .env 자동 로드 (로컬 .env → sns_analyzer .env 순으로 탐색)
def _load_env():
    candidates = [
        os.path.join(os.path.dirname(__file__), ".env"),
        os.path.join(os.path.dirname(__file__), "..", "sns_analyzer", ".env"),
        r"E:\sns_analyzer\.env",
    ]
    for path in candidates:
        if os.path.exists(path):
            try:
                from dotenv import dotenv_values
                for k, v in dotenv_values(path).items():
                    if k not in os.environ and v:
                        os.environ[k] = v
                return
            except Exception:
                pass

_load_env()

if not os.environ.get("ANTHROPIC_API_KEY"):
    print("⚠️  ANTHROPIC_API_KEY 환경변수가 없습니다.")
    print("   Codespaces: github.com/settings/codespaces 에서 Secret 등록 후 Codespace 재시작")
    print("   로컬: .env 파일에 ANTHROPIC_API_KEY=sk-ant-... 추가")

# ── History DB ────────────────────────────────────────────────────────────────
DB_PATH = os.path.join(os.path.dirname(__file__), "history.db")

def _init_db():
    with sqlite3.connect(DB_PATH) as c:
        c.execute("""CREATE TABLE IF NOT EXISTS analyses (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            idea         TEXT NOT NULL,
            scope        TEXT,
            industry     TEXT,
            stage        TEXT,
            depth        TEXT,
            currency     TEXT,
            full_text    TEXT,
            tokens_input  INTEGER DEFAULT 0,
            tokens_output INTEGER DEFAULT 0,
            tokens_cached INTEGER DEFAULT 0,
            created_at   TEXT DEFAULT (datetime('now','localtime'))
        )""")
        c.commit()

_init_db()

from fastapi import FastAPI, HTTPException
from fastapi.responses import StreamingResponse, FileResponse, Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import List

app = FastAPI(title="사업구상 분석", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

SYSTEM_PROMPT = """당신은 YC·Sequoia·a16z 출신 최상급 벤처 투자자이자 맥킨지·골드만삭스 수준의 사업 전략 컨설턴트, 시장 분석가, 브랜드 전략가이다.
10억 달러 이상의 엑싯을 목표로 하는 투자자의 냉정한 시각으로 분석한다.

사용자가 사업 아이디어를 입력하면 아래 항목을 체계적으로 분석한다.

분석 목적:
- 사업 타당성 검토
- 시장 규모 추정
- 경쟁 우위 도출
- 수익 가능성 평가
- 브랜딩 방향 설정
- 실행 전략 제안

항상 실제 시장 기준으로 현실적이고 비판적으로 분석한다.
근거 없는 낙관론은 금지한다.
가능하면 수치 추정과 논리 근거를 함께 제시한다.

다음 형식을 반드시 따른다.

---

# 1. 사업 아이디어 요약

- 사업 개념:
- 핵심 고객:
- 해결하려는 문제:
- 제공 가치:
- 예상 수익 모델:
- 차별화 포인트:

---

# 2. 시장 분석

## TAM (Total Addressable Market)
- 전체 잠재 시장 규모
- 글로벌 기준 / 국내 기준 분리
- 가능한 경우 금액 규모 추정
- 시장 성장률 포함

## SAM (Serviceable Available Market)
- 실제 접근 가능한 시장
- 타겟 고객 범위
- 초기 집중 세그먼트

## SOM (Serviceable Obtainable Market)
- 현실적으로 확보 가능한 시장 점유율
- 1년 / 3년 / 5년 예상

시장 규모 계산 논리를 반드시 설명한다.

---

# 3. 고객 분석

## 핵심 고객 페르소나
- 연령
- 직업
- 구매력
- 행동 패턴
- 주요 Pain Point

## 고객이 현재 사용하는 대안
- 기존 서비스
- 수작업 방식
- 경쟁 제품

## 구매 동기
- 비용 절감
- 시간 절약
- 수익 증가
- 편의성
- 감정적 만족

---

# 4. 경쟁사 분석

최소 5개 경쟁사 분석.

표 형식으로:
- 경쟁사명
- 핵심 기능
- 가격 정책
- 강점
- 약점
- 시장 포지션

그리고 아래 추가 분석 포함:
- 시장 리더
- 빠르게 성장 중인 플레이어
- 아직 해결되지 않은 시장 공백
- 사용자가 불만을 느끼는 부분

---

# 5. SWOT 분석

## Strengths
내부 강점

## Weaknesses
내부 약점

## Opportunities
시장 기회

## Threats
외부 위협

현실적으로 작성한다.
투자자 관점으로 냉정하게 분석한다.

---

# 6. 수익 모델 분석

가능한 BM 제안:
- SaaS
- 구독형
- 광고
- 중개 수수료
- 라이선스
- 프리미엄 모델
- API 판매
- 컨설팅 연계
- 데이터 판매

각 모델의 장단점과 적합도를 설명한다.

---

# 7. 실행 난이도 분석

다음을 10점 만점으로 평가:
- 개발 난이도
- 시장 진입 난이도
- 마케팅 난이도
- 자본 필요도
- 경쟁 강도

그리고 이유 설명.

---

# 8. MVP 전략

초기 MVP에 반드시 필요한 기능만 추려라.

포함:
- 핵심 기능
- 제거 가능한 기능
- 1개월 MVP 버전
- 사용자 검증 방법
- 빠른 시장 테스트 방법

Lean Startup 방식으로 설명한다.

---

# 9. 브랜드 전략

## 브랜드 방향성
- 고급형 / 대중형 / 혁신형 등

## 추천 브랜드 네이밍
최소 15개 제안.

조건:
- 짧고 기억하기 쉬울 것
- 글로벌 발음 가능
- 도메인 브랜드화 가능성 고려
- 너무 흔한 단어 조합 지양

각 이름마다:
- 의미
- 브랜드 이미지
- 추천 이유

포함.

---

# 10. 마케팅 전략

초기 고객 확보 전략:
- SEO
- SNS
- 바이럴
- 커뮤니티
- B2B 세일즈
- 광고
- 콘텐츠 마케팅
- 인플루언서
- 파트너십

채널 우선순위 제시.

---

# 11. 투자 매력도 평가

다음을 10점 만점으로 평가하고, 각 항목마다 점수 근거와 투자 유치 후 18개월 이내 개선 전략을 제시한다:
- 시장성
- 확장성
- 수익성
- 방어력
- 글로벌 가능성
- AI 대체 위험
- 투자 매력도

총평 제공 후, 반드시 마지막 두 줄에 아래 형식으로 출력:
[INVESTMENT: 시장성=숫자, 확장성=숫자, 수익성=숫자, 방어력=숫자, 글로벌=숫자, AI위험=숫자, 매력도=숫자]
[INVEST_AFTER: 시장성=숫자, 확장성=숫자, 수익성=숫자, 방어력=숫자, 글로벌=숫자, AI위험=숫자, 매력도=숫자]
(INVEST_AFTER는 투자 유치 성공 후 18~24개월 시점 예상 점수, 현재보다 반드시 높아야 함)

---

# 12. 최종 결론

아래 항목 포함:
- 이 사업이 성공할 가능성이 높은 이유
- 실패 가능성이 높은 이유
- 가장 중요한 성공 변수
- 지금 시작해도 되는지 여부
- 추천 액션 플랜 TOP 5

---

---

# 13. 수익 시뮬레이션

연도별 재무 모델을 3가지 시나리오로 제시한다.
모든 수치는 근거 있는 추정값이어야 한다.

핵심 가정 먼저 명시:
- 초기 고객 획득 비용 (CAC)
- 고객 생애 가치 (LTV) 및 LTV:CAC 비율
- 월 번 레이트 (Burn Rate)
- 손익분기점 (BEP) 도달 시점

## 시나리오별 5개년 재무 추정

표 형식으로 보수적 / 기본 / 낙관적 시나리오 각각 제시.

포함 항목:
- 연간 MAU (또는 계약 고객 수)
- ARPU / 객단가
- 총 매출
- 주요 비용 항목 (인건비, 마케팅, 인프라)
- EBITDA / 영업이익
- 누적 투자 소진액

## 핵심 재무 지표
- Gross Margin %
- CAC 회수 기간
- Net Revenue Retention (해당 시)
- 연간 성장률 (YoY)

---

# 14. 글로벌 확장 전략

## Phase 1: 국내 시장 장악 (0~18개월)
- 핵심 KPI 목표
- 지역 집중 전략
- 레퍼런스 확보 방법

## Phase 2: 1차 해외 시장 진출 (18~36개월)

우선순위 시장 선정 (최소 3개국):
| 국가 | 선정 이유 | 시장 규모 | 진입 전략 | 리스크 |
|------|-----------|-----------|-----------|--------|

진입 방식:
- 직접 진출 vs 현지 파트너십 vs 인수합병
- 현지화 필요 요소 (언어, 규제, 문화)
- 예상 비용 및 타임라인

## Phase 3: 글로벌 스케일 (36개월+)
- 멀티 리전 운영 전략
- 글로벌 브랜드 포지셔닝

## 글로벌화 핵심 장벽
- 규제 및 법률 이슈
- 문화적 적응 필요 요소
- 현지 강자와의 경쟁
- 기술/인프라 현지화 비용

---

# 15. VC 투자자 심사 관점

YC Demo Day / Sequoia Pitch 기준으로 이 사업을 심사한다.
투자자가 이 피치를 듣고 즉시 수표를 쓰고 싶어지도록 작성한다.

## 투자 명제 (Investment Thesis)
한 문장으로 핵심 투자 논리를 작성한다.
형식: "[독보적 포지션]을 통해 [시장 규모]에서 [달성 목표]를 실현한다. [지금 투자해야 하는 결정적 이유]"

## Why Now? — 지금이 최적의 투자 타이밍인 이유
- 기술 트렌드 (AI, 플랫폼 변화 등)
- 규제/정책 변화
- 소비자 행동 변화
- 경쟁 공백이 지금 열려있는 이유

## 10x 수익 가능성
구체적 수치로 투자 ROI 시나리오 제시 (3년, 5년 기준 멀티플)

## 기술적·구조적 해자 (Moat)
경쟁자가 18개월 안에 복제할 수 없는 진입 장벽 (데이터, 네트워크 효과, 특허, 브랜드)

## 투자 단계별 마일스톤

| 단계 | 조달 금액 | 핵심 마일스톤 | 달성 시점 | 자금 사용처 |
|------|-----------|--------------|----------|------------|
| Pre-seed | | | | |
| Seed | | | | |
| Series A | | | | |
| Series B | | | | |

## 팀 요건
이 사업을 반드시 이길 팀의 조건과 채용 전략

## VC 원라이너 평가
"이 딜에 투자할 것인가?" — YES/NO 명확히 + 한 줄 근거
투자 거절 시 극복 방법도 제시

---

출력 스타일:
- 맥킨지 전략 보고서 + YC 투자 심사 수준으로 작성
- 핵심 위주, 불필요한 미사여구 금지
- 표 적극 활용
- 실제 유사 기업 사례 반드시 언급 (성공/실패 모두)
- 모든 수치는 근거와 함께 제시
- 중요한 부분은 bullet point 사용
- 비판적 시각 유지: 좋은 점만큼 문제점도 명확히"""


# ── Options ───────────────────────────────────────────────────────────────────

SCOPE_NAMES = {
    "global": "전체 글로벌", "korea": "대한민국", "usa": "미국",
    "japan": "일본", "china": "중국", "southeast_asia": "동남아시아",
    "europe": "유럽", "middle_east": "중동·아프리카", "latam": "남미",
}
STAGE_NAMES = {
    "idea": "아이디어 단계 (검증 전)",
    "prototype": "MVP/프로토타입 완성",
    "seed": "초기 고객 확보 중",
    "growth": "성장 단계 (PMF 확인)",
}
DEPTH_TOKENS = {"quick": 5000, "standard": 8000, "deep": 16000}


LANG_INSTRUCTION = {
    "en": "Write ALL analysis content in English. All section headings, bullet points, tables, and conclusions must be in English.",
    "ja": "分析内容をすべて日本語で記述してください。見出し、箇条書き、表、結論をすべて日本語で記述すること。",
    "zh": "请将所有分析内容用中文撰写。所有标题、要点、表格和结论均须使用中文。",
}

class AnalyzeRequest(BaseModel):
    idea: str
    scope: List[str] = ["global"]
    industry: str = ""
    stage: str = "idea"
    depth: str = "standard"
    currency: str = "KRW"
    lang: str = "ko"


def build_context(body: AnalyzeRequest) -> str:
    scope_text = " / ".join(SCOPE_NAMES.get(s, s) for s in (body.scope or ["global"]))
    lines = [
        "\n\n[분석 컨텍스트 — 반드시 이 조건으로 분석]",
        f"- 목표 시장: {scope_text}",
        f"- 사업 단계: {STAGE_NAMES.get(body.stage, body.stage)}",
        f"- 통화 단위: {body.currency} (모든 금액 표기에 사용)",
    ]
    if body.industry:
        lines.append(f"- 산업 분야: {body.industry} (해당 산업 특화 분석 반영)")
    if body.depth == "deep":
        lines.append("- 분석 깊이: 심층 분석. 각 섹션 최대한 상세하게 작성.")
    elif body.depth == "quick":
        lines.append("- 분석 깊이: 빠른 검토. 핵심 인사이트만 간결하게.")

    lines += [
        "",
        "[필수 차트 데이터 — 반드시 해당 섹션 마지막에 정확히 출력]",
        "섹션 2 (시장 분석) 마지막 줄:",
        f"[MARKET: TAM=숫자, SAM=숫자, SOM=숫자, unit={'B USD' if body.currency=='USD' else '억엔' if body.currency=='JPY' else 'B EUR' if body.currency=='EUR' else '억원'}]",
        f"  예시(USD): [MARKET: TAM=850, SAM=120, SOM=12, unit=B USD]  (십억달러 단위 순수 정수)",
        "섹션 7 (실행 난이도) 마지막 줄:",
        "[DIFFICULTY: 개발=숫자, 시장진입=숫자, 마케팅=숫자, 자본=숫자, 경쟁=숫자]",
        "  예시: [DIFFICULTY: 개발=7, 시장진입=6, 마케팅=5, 자본=8, 경쟁=7]",
        "섹션 11 (투자 매력도) 마지막 줄:",
        "[INVESTMENT: 시장성=숫자, 확장성=숫자, 수익성=숫자, 방어력=숫자, 글로벌=숫자, AI위험=숫자, 매력도=숫자]",
        "섹션 10 (마케팅 전략) 마지막 줄:",
        "[CHANNELS: SEO=숫자, SNS=숫자, 바이럴=숫자, 커뮤니티=숫자, B2B=숫자, 광고=숫자, 콘텐츠=숫자, 인플루언서=숫자, 파트너십=숫자]",
        "  예시: [CHANNELS: SEO=8, SNS=9, 바이럴=7, 커뮤니티=6, B2B=5, 광고=4, 콘텐츠=8, 인플루언서=5, 파트너십=6]  (중요도 1-10)",
        "섹션 13 (수익 시뮬레이션) 마지막 줄:",
        "[REVENUE: cons=숫자:숫자:숫자:숫자:숫자, base=숫자:숫자:숫자:숫자:숫자, opt=숫자:숫자:숫자:숫자:숫자, unit=억원]",
        "  예시: [REVENUE: cons=1:3:8:15:25, base=2:8:20:40:70, opt=5:20:50:100:180, unit=억원]  (연도1~5 매출, 순수 정수)",
    ]
    return "\n".join(lines)


# ── DOC Export ────────────────────────────────────────────────────────────────

def md_to_docx(text: str, idea: str, theme: str = "dark") -> bytes:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    is_light = theme == "light"
    COLOR_TEXT   = RGBColor(0x1e, 0x29, 0x3b) if is_light else RGBColor(0xe2, 0xe8, 0xf0)
    COLOR_MUTED  = RGBColor(0x47, 0x55, 0x69) if is_light else RGBColor(0x6b, 0x72, 0x80)
    COLOR_ACCENT = RGBColor(0x37, 0x30, 0xa3) if is_light else RGBColor(0x81, 0x8c, 0xf8)
    BG_HEX       = "F1F5F9" if is_light else "07070F"

    doc = Document()

    # Set page background color
    bg = doc.settings.element
    bg_elem = OxmlElement('w:displayBackgroundShape')
    bg.append(bg_elem)
    body_bg = doc.element.body
    sectPr = body_bg.find(qn('w:sectPr'))
    if sectPr is None:
        sectPr = OxmlElement('w:sectPr')
        body_bg.append(sectPr)
    bg_color = OxmlElement('w:background')
    bg_color.set(qn('w:color'), BG_HEX)
    doc.element.insert(0, bg_color)

    for sec in doc.sections:
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
        sec.top_margin = Cm(2.0)
        sec.bottom_margin = Cm(2.0)

    h = doc.add_heading("사업구상 분석 리포트", 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if h.runs:
        h.runs[0].font.color.rgb = COLOR_ACCENT
    sub = doc.add_paragraph(f"분석 대상: {idea}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if sub.runs:
        sub.runs[0].font.size = Pt(12)
        sub.runs[0].font.color.rgb = COLOR_MUTED
    doc.add_paragraph()

    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]

        if re.match(r"^\[(MARKET|DIFFICULTY|INVESTMENT):", line):
            i += 1
            continue

        if line.startswith("| "):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s\-|:]+\|$", lines[i]):
                    rows.append([c.strip() for c in lines[i].split("|")[1:-1]])
                i += 1
            if rows:
                ncols = max(len(r) for r in rows)
                tbl = doc.add_table(rows=len(rows), cols=ncols)
                tbl.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci in range(ncols):
                        cell = tbl.rows[ri].cells[ci]
                        cell.text = row[ci] if ci < len(row) else ""
                        if ri == 0 and cell.paragraphs[0].runs:
                            cell.paragraphs[0].runs[0].bold = True
                doc.add_paragraph()
            continue

        def set_run_color(run, color):
            run.font.color.rgb = color

        if line.startswith("# "):
            p = doc.add_heading(line[2:], level=1)
            for r in p.runs: set_run_color(r, COLOR_ACCENT)
        elif line.startswith("## "):
            p = doc.add_heading(line[3:], level=2)
            for r in p.runs: set_run_color(r, COLOR_ACCENT)
        elif line.startswith("### "):
            p = doc.add_heading(line[4:], level=3)
            for r in p.runs: set_run_color(r, COLOR_MUTED)
        elif line == "---":
            doc.add_paragraph()
        elif re.match(r"^[-*] ", line):
            p = doc.add_paragraph(line[2:], style="List Bullet")
            for r in p.runs: set_run_color(r, COLOR_TEXT)
        elif line.strip():
            para = doc.add_paragraph()
            for part in re.split(r"(\*\*.+?\*\*)", line):
                if part.startswith("**") and part.endswith("**"):
                    run = para.add_run(part[2:-2])
                    run.bold = True
                    set_run_color(run, COLOR_TEXT)
                elif part:
                    run = para.add_run(part)
                    set_run_color(run, COLOR_TEXT)
        i += 1

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def md_to_pptx(text: str, idea: str, theme: str = "dark") -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    import re as _re

    is_light = theme == "light"

    BG      = RGBColor(0xf1,0xf5,0xf9) if is_light else RGBColor(0x0c,0x0c,0x1a)
    TEXT    = RGBColor(0x1e,0x29,0x3b) if is_light else RGBColor(0xff,0xff,0xff)
    TEXT2   = RGBColor(0x33,0x41,0x55) if is_light else RGBColor(0xd1,0xd5,0xdb)
    MUTED   = RGBColor(0x64,0x74,0x8b) if is_light else RGBColor(0x9c,0xa3,0xaf)
    CARD_BG = RGBColor(0xe2,0xe8,0xf0) if is_light else RGBColor(0x10,0x10,0x24)
    DIVIDER = RGBColor(0xcb,0xd5,0xe1) if is_light else RGBColor(0x1e,0x1e,0x32)
    BRAND   = RGBColor(0x63,0x66,0xf1)

    ACCENTS = [
        RGBColor(0x63,0x66,0xf1), RGBColor(0x22,0xd3,0xee), RGBColor(0xa7,0x8b,0xfa),
        RGBColor(0xf8,0x71,0x71), RGBColor(0xfb,0xbf,0x24), RGBColor(0x34,0xd3,0x99),
        RGBColor(0xf8,0x71,0x71), RGBColor(0x63,0x66,0xf1), RGBColor(0xa7,0x8b,0xfa),
        RGBColor(0x22,0xd3,0xee), RGBColor(0x34,0xd3,0x99), RGBColor(0xfb,0xbf,0x24),
    ]

    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)
    BLANK = prs.slide_layouts[6]
    W, H = 13.33, 7.5

    def add_rect(slide, l, t, w, h, fill, line_col=None):
        s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
        s.fill.solid(); s.fill.fore_color.rgb = fill
        if line_col:
            s.line.fill.solid(); s.line.fill.fore_color.rgb = line_col
            s.line.width = Pt(0.75)
        else:
            s.line.fill.background()
        return s

    def add_tb(slide, txt, l, t, w, h, sz=11, bold=False, col=None, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
        box.word_wrap = True
        tf = box.text_frame; tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = align
        r = p.add_run(); r.text = txt
        r.font.size = Pt(sz); r.font.bold = bold
        r.font.color.rgb = col or MUTED
        return box

    def add_bg(slide):
        add_rect(slide, 0, 0, W, H, BG)

    # ── Title slide ──
    s0 = prs.slides.add_slide(BLANK); add_bg(s0)
    add_rect(s0, 0, 0, W, 0.07, BRAND)
    add_rect(s0, W/2 - 0.5, 3.92, 1.0, 0.04, BRAND)
    add_tb(s0, 'Pitch Deck', 1, 1.9, W-2, 0.42, sz=13, bold=True, col=ACCENTS[0], align=PP_ALIGN.CENTER)
    add_tb(s0, idea, 1, 2.45, W-2, 1.3, sz=38, bold=True, col=TEXT, align=PP_ALIGN.CENTER)
    add_tb(s0, 'Claude Opus  McKinsey/YC Level Analysis', 1, 4.1, W-2, 0.42, sz=11, col=MUTED, align=PP_ALIGN.CENTER)

    # ── Parse sections ──
    SKIP = _re.compile(r'^\[(MARKET|DIFFICULTY|INVESTMENT|REVENUE|CHANNELS):')
    secs = []
    cur = None
    for line in text.split('\n'):
        m = _re.match(r'^# (\d+)\.\s*(.+)', line)
        if m:
            if cur: secs.append(cur)
            cur = {'num': int(m.group(1)), 'title': m.group(2).strip(), 'lines': []}
        elif cur:
            cur['lines'].append(line)
    if cur: secs.append(cur)

    def clean_md(s):
        s = _re.sub(r'\*\*(.+?)\*\*', r'\1', s)
        s = _re.sub(r'\*(.+?)\*', r'\1', s)
        return s.strip()

    def parse_content(lines):
        groups, kvs, bullets, grp = [], [], [], None
        for ln in lines:
            stripped = ln.strip()
            if SKIP.match(stripped) or not stripped or stripped == '---':
                continue
            # Catch ALL table lines: |...|, including |---|---|--- separators
            if stripped.startswith('|'):
                if _re.match(r'^\|[\s\-|:]+\|', stripped):  # separator row
                    continue
                cells = [clean_md(c) for c in stripped.split('|')[1:-1] if c.strip()]
                if cells:
                    (grp['items'] if grp is not None else bullets).append(
                        '  |  '.join(cells)[:85])
                continue
            if stripped.startswith('## ') or stripped.startswith('### '):
                grp = {'heading': clean_md(stripped.lstrip('#')), 'items': []}
                groups.append(grp)
            elif _re.match(r'^[-*] ', stripped):
                item = clean_md(stripped[2:])
                (grp['items'] if grp is not None else bullets).append(item)
            elif ':' in stripped and not stripped.startswith('#'):
                k, _, v = stripped.partition(':')
                k = clean_md(k.lstrip('-')); v = clean_md(v)
                if k and v and len(k) < 50:
                    kvs.append({'k': k, 'v': v})
                elif grp is not None:
                    grp['items'].append(clean_md(stripped)[:85])
                else:
                    bullets.append(clean_md(stripped)[:85])
            else:
                t = clean_md(stripped)
                if t and not t.startswith('#'):
                    (grp['items'] if grp is not None else bullets).append(t)
        return groups, kvs, bullets

    # Layout constants
    BYX = 1.72                    # body area start y
    CX1, CW1 = 0.45, 6.3         # left column x, width
    CX2, CW2 = 7.38, 5.5         # right column x, width

    def render_bullets(slide, items, x, y, w, ac, maxn=13):
        cy = y
        for it in (items or [])[:maxn]:
            if cy > H - 0.35: break
            add_rect(slide, x, cy+0.072, 0.065, 0.065, ac)
            add_tb(slide, (it[:90]+'...' if len(it)>90 else it), x+0.14, cy, w-0.14, 0.31, sz=10.5, col=TEXT2)
            cy += 0.33

    def render_groups(slide, groups, x, y, w, ac, maxitems=4):
        cy = y
        for g in (groups or []):
            if cy > H - 0.5: break
            n = min(len(g.get('items') or []), maxitems)
            ch = 0.33 + n * 0.27 + 0.1
            add_rect(slide, x, cy, w, ch, CARD_BG, DIVIDER)
            if g.get('heading'):
                add_tb(slide, g['heading'][:55], x+0.14, cy+0.07, w-0.28, 0.22, sz=8.5, bold=True, col=ac)
            iy = cy + 0.3
            for it in (g.get('items') or [])[:maxitems]:
                if iy > H - 0.3: break
                add_rect(slide, x+0.14, iy+0.075, 0.055, 0.055, ac)
                add_tb(slide, (it[:78]+'...' if len(it)>78 else it), x+0.25, iy, w-0.4, 0.26, sz=10, col=TEXT2)
                iy += 0.27
            cy += ch + 0.1

    def render_kv(slide, kvs, x, y, w, ac, maxn=8):
        cy = y
        for kv in (kvs or [])[:maxn]:
            if cy > H - 0.4: break
            add_rect(slide, x, cy, w, 0.36, CARD_BG, DIVIDER)
            add_tb(slide, kv['k'][:38], x+0.13, cy+0.05, w*0.44, 0.27, sz=9.5, col=MUTED)
            v = (kv['v'][:46]+'...' if len(kv['v'])>46 else kv['v'])
            add_tb(slide, v, x+w*0.46, cy+0.05, w*0.52, 0.27, sz=10, bold=True, col=ac)
            cy += 0.42

    def col_divider(slide, x, y, bh):
        add_rect(slide, x, y, 0.02, bh, DIVIDER)

    for i, sec in enumerate(secs):
        ac = ACCENTS[i % len(ACCENTS)]
        slide = prs.slides.add_slide(BLANK); add_bg(slide)
        groups, kvs, bullets = parse_content(sec['lines'])
        bh = H - BYX - 0.15

        # Accent top bar
        add_rect(slide, 0, 0, W, 0.07, ac)
        # Header / body separator
        add_rect(slide, 0, 1.65, W, 0.02, DIVIDER)
        # Tag label (small caps accent)
        add_tb(slide, f"SECTION {sec['num']}", 0.5, 0.19, 4, 0.3, sz=9, bold=True, col=ac)
        # Section title
        add_tb(slide, sec['title'], 0.5, 0.5, W-1.0, 1.12, sz=22, bold=True, col=TEXT)

        if len(groups) >= 3:
            render_groups(slide, groups[:2], CX1, BYX, CW1, ac, 4)
            col_divider(slide, CX2-0.08, BYX, bh)
            render_groups(slide, groups[2:4], CX2, BYX, CW2, ac, 4)
        elif groups and kvs:
            render_groups(slide, groups[:2], CX1, BYX, CW1, ac, 4)
            col_divider(slide, CX2-0.08, BYX, bh)
            render_kv(slide, kvs[:7], CX2, BYX, CW2, ac)
        elif groups and bullets:
            render_groups(slide, groups[:2], CX1, BYX, CW1, ac, 5)
            col_divider(slide, CX2-0.08, BYX, bh)
            render_bullets(slide, bullets[:11], CX2, BYX, CW2, ac)
        elif groups:
            render_groups(slide, groups, CX1, BYX, W-0.9, ac, 4)
        elif kvs and bullets:
            render_kv(slide, kvs[:7], CX1, BYX, CW1, ac)
            col_divider(slide, CX2-0.08, BYX, bh)
            render_bullets(slide, bullets[:11], CX2, BYX, CW2, ac)
        elif kvs:
            render_kv(slide, kvs[:9], CX1, BYX, W-0.9, ac)
        else:
            render_bullets(slide, bullets[:14], CX1, BYX, W-0.9, ac)

    buf = io.BytesIO()
    prs.save(buf); buf.seek(0)
    return buf.read()


class ExportRequest(BaseModel):
    text: str
    idea: str = "사업 분석"
    theme: str = "dark"


@app.post("/api/export")
async def export_doc(body: ExportRequest):
    try:
        docx_bytes = await asyncio.to_thread(md_to_docx, body.text, body.idea, body.theme)
        return Response(
            content=docx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": 'attachment; filename="biz-analysis.docx"'},
        )
    except ImportError:
        raise HTTPException(500, "python-docx 패키지가 필요합니다.")
    except Exception as e:
        raise HTTPException(500, str(e))


@app.post("/api/export/ppt")
async def export_ppt(body: ExportRequest):
    try:
        pptx_bytes = await asyncio.to_thread(md_to_pptx, body.text, body.idea, body.theme)
        return Response(
            content=pptx_bytes,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers={"Content-Disposition": 'attachment; filename="biz-analysis.pptx"'},
        )
    except ImportError:
        raise HTTPException(500, "python-pptx 패키지가 필요합니다.")
    except Exception as e:
        raise HTTPException(500, str(e))


# ── History endpoints ─────────────────────────────────────────────────────────

class SaveRequest(BaseModel):
    idea: str
    scope: List[str] = ["global"]
    industry: str = ""
    stage: str = "idea"
    depth: str = "standard"
    currency: str = "USD"
    full_text: str
    tokens_input: int = 0
    tokens_output: int = 0
    tokens_cached: int = 0


@app.post("/api/history")
async def save_history(body: SaveRequest):
    def _save():
        with sqlite3.connect(DB_PATH) as c:
            c.execute(
                "INSERT INTO analyses (idea,scope,industry,stage,depth,currency,full_text,tokens_input,tokens_output,tokens_cached) VALUES (?,?,?,?,?,?,?,?,?,?)",
                (body.idea, json.dumps(body.scope), body.industry, body.stage, body.depth,
                 body.currency, body.full_text, body.tokens_input, body.tokens_output, body.tokens_cached),
            )
            c.commit()
    await asyncio.to_thread(_save)
    return {"ok": True}


@app.get("/api/history")
async def list_history():
    def _list():
        with sqlite3.connect(DB_PATH) as c:
            c.row_factory = sqlite3.Row
            rows = c.execute(
                "SELECT id,idea,scope,industry,stage,depth,currency,tokens_input,tokens_output,tokens_cached,created_at FROM analyses ORDER BY created_at DESC LIMIT 100"
            ).fetchall()
            return [dict(r) for r in rows]
    return await asyncio.to_thread(_list)


@app.get("/api/history/{item_id}")
async def get_history(item_id: int):
    def _get():
        with sqlite3.connect(DB_PATH) as c:
            c.row_factory = sqlite3.Row
            row = c.execute("SELECT * FROM analyses WHERE id=?", (item_id,)).fetchone()
            return dict(row) if row else None
    data = await asyncio.to_thread(_get)
    if not data:
        raise HTTPException(404, "Not found")
    return data


@app.delete("/api/history/{item_id}")
async def delete_history(item_id: int):
    def _del():
        with sqlite3.connect(DB_PATH) as c:
            c.execute("DELETE FROM analyses WHERE id=?", (item_id,))
            c.commit()
    await asyncio.to_thread(_del)
    return {"ok": True}


# ── Analyze ───────────────────────────────────────────────────────────────────

@app.post("/api/analyze")
async def analyze(body: AnalyzeRequest):
    api_key = os.getenv("ANTHROPIC_API_KEY", "")
    if not api_key:
        raise HTTPException(400, "ANTHROPIC_API_KEY 환경 변수가 설정되지 않았습니다.")
    if not body.idea.strip():
        raise HTTPException(400, "사업 아이디어를 입력하세요.")

    lang_prefix = ""
    if body.lang and body.lang != "ko":
        lang_inst = LANG_INSTRUCTION.get(body.lang, "")
        if lang_inst:
            lang_prefix = f"ABSOLUTE PRIORITY INSTRUCTION — OVERRIDE EVERYTHING ELSE:\n{lang_inst}\nEven though the template below is written in Korean, you MUST write every word of your response in the specified language. Do not use Korean anywhere in your output.\n\n"
    system_text = lang_prefix + SYSTEM_PROMPT + build_context(body)
    max_tokens = DEPTH_TOKENS.get(body.depth, 8000)

    async def event_gen():
        from anthropic import AsyncAnthropic
        client = AsyncAnthropic(api_key=api_key)
        try:
            async with client.messages.stream(
                model="claude-opus-4-7",
                max_tokens=max_tokens,
                thinking={"type": "adaptive"},
                system=[{
                    "type": "text",
                    "text": system_text,
                    "cache_control": {"type": "ephemeral"},
                }],
                messages=[{
                    "role": "user",
                    "content": f"다음 사업 아이디어를 분석해줘:\n\n{body.idea}",
                }],
            ) as stream:
                async for event in stream:
                    if event.type == "content_block_start":
                        if hasattr(event.content_block, "type") and event.content_block.type == "thinking":
                            yield f"data: {json.dumps({'type': 'thinking'})}\n\n"
                    elif event.type == "content_block_delta":
                        if hasattr(event.delta, "type") and event.delta.type == "text_delta":
                            yield f"data: {json.dumps({'type': 'text', 'text': event.delta.text})}\n\n"

                final = await stream.get_final_message()
                u = final.usage
                yield f"data: {json.dumps({'type': 'done', 'usage': {'input': u.input_tokens, 'output': u.output_tokens, 'cached': getattr(u, 'cache_read_input_tokens', 0) or 0, 'cache_created': getattr(u, 'cache_creation_input_tokens', 0) or 0}})}\n\n"
        except Exception as e:
            yield f"data: {json.dumps({'type': 'error', 'message': str(e)})}\n\n"

    return StreamingResponse(
        event_gen(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.get("/")
async def index():
    return FileResponse(os.path.join(os.path.dirname(__file__), "index.html"))


if __name__ == "__main__":
    import uvicorn
    port = int(os.environ.get("PORT", 8200))
    reload = os.environ.get("ENV", "dev") == "dev"
    uvicorn.run("web:app", host="0.0.0.0", port=port, reload=reload)
